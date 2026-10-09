"""Creates Production_Tracker_Template.xlsx (upload to Google Drive and open as a Google Sheet)."""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from build import LOG_COLUMNS, TRACKER_COLUMNS

HEADER = PatternFill('solid', fgColor='1F3A5F')
WIDTHS = {'Item ID': 14, 'Description': 34, 'Category': 14, 'Location / Line': 16, 'Unit': 8,
          'Planned Qty': 12, 'Actual Qty': 12, 'Rejected Qty': 12, 'Remaining Qty': 13, '% Complete': 11,
          'Planned Finish': 14, 'Status': 16, 'Remarks': 30, 'Last Updated': 17, 'Source File': 30}
STATUS_COLORS = {'Complete': 'D9EAD3', 'In Progress': 'DCE6F1', 'Behind Schedule': 'F4CCCC',
                 'On Hold': 'FFF2CC', 'Not Started': 'EEEEEE', 'Cancelled': 'D9D9D9'}


def header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = widths.get(c, 16)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(cols)).column_letter}1'


wb = Workbook()
ws = wb.active
ws.title = 'Tracker'
header(ws, TRACKER_COLUMNS, WIDTHS)
col = {c: ws.cell(row=1, column=i).column_letter for i, c in enumerate(TRACKER_COLUMNS, 1)}
last = col['Source File']
for r in range(2, 1001):
    ws[f"{col['% Complete']}{r}"].number_format = '0%'
    ws[f"{col['Planned Finish']}{r}"].number_format = 'yyyy-mm-dd'
dv = DataValidation(type='list', formula1='"' + ','.join(STATUS_COLORS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{col['Status']}2:{col['Status']}1000")
for status, color in STATUS_COLORS.items():
    ws.conditional_formatting.add(f'A2:{last}1000', FormulaRule(
        formula=[f'${col["Status"]}2="{status}"'], fill=PatternFill('solid', fgColor=color)))

log = wb.create_sheet('Update Log')
header(log, LOG_COLUMNS, {'Timestamp': 17, 'Source File': 34, 'File ID': 20, 'Result': 44})

wb.save('Production_Tracker_Template.xlsx')
print('wrote Production_Tracker_Template.xlsx')
